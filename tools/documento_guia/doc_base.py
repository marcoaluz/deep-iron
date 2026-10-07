"""Utilitários de formatação do documento DEEP IRON (python-docx)."""
import os
from docx import Document
from docx.enum.section import WD_ORIENT
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor

FERRUGEM = RGBColor(0x8A, 0x3B, 0x12)
CARVAO = RGBColor(0x2B, 0x22, 0x1C)
CINZA = RGBColor(0x55, 0x50, 0x4A)
ARTE = r"D:\DEV\deep-iron\docs\arte"

doc = Document()
sec = doc.sections[0]
sec.page_width, sec.page_height = Cm(21), Cm(29.7)
for m in ("left_margin", "right_margin"):
    setattr(sec, m, Cm(2.0))
sec.top_margin, sec.bottom_margin = Cm(2.0), Cm(1.8)

st = doc.styles["Normal"]
st.font.name = "Calibri"
st.font.size = Pt(10.5)
st.element.rPr.rFonts.set(qn("w:eastAsia"), "Calibri")
st.paragraph_format.space_after = Pt(4)
st.paragraph_format.line_spacing = 1.12
for nivel, tam, cor in ((1, 18, FERRUGEM), (2, 14, CARVAO), (3, 12, FERRUGEM)):
    h = doc.styles["Heading %d" % nivel]
    h.font.name = "Calibri"
    h.font.size = Pt(tam)
    h.font.bold = True
    h.font.color.rgb = cor
    h.element.rPr.rFonts.set(qn("w:eastAsia"), "Calibri")
    h.paragraph_format.space_before = Pt(14 if nivel == 1 else 10)
    h.paragraph_format.space_after = Pt(5)
    h.paragraph_format.keep_with_next = True


def H(texto, nivel=1, nova_pagina=False):
    if nova_pagina:
        doc.add_paragraph().add_run().add_break(WD_BREAK.PAGE)
    return doc.add_heading(texto, level=nivel)


def _runs(p, texto):
    """**negrito** dentro do texto."""
    partes = texto.split("**")
    for i, t in enumerate(partes):
        if t:
            r = p.add_run(t)
            r.bold = i % 2 == 1
    return p


def P(texto, italico=False, cor=None, tam=None, alinhar=None):
    p = _runs(doc.add_paragraph(), texto)
    for r in p.runs:
        if italico:
            r.italic = True
        if cor:
            r.font.color.rgb = cor
        if tam:
            r.font.size = Pt(tam)
    if alinhar == "centro":
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    return p


def B(itens, numerada=False):
    estilo = "List Number" if numerada else "List Bullet"
    for it in itens:
        if isinstance(it, (list, tuple)):  # (texto, [subitens])
            _runs(doc.add_paragraph(style=estilo), it[0])
            for sub in it[1]:
                p = _runs(doc.add_paragraph(style="List Bullet 2"), sub)
        else:
            _runs(doc.add_paragraph(style=estilo), it)


def _sombra(celula, hexcor):
    tcPr = celula._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:color"), "auto")
    shd.set(qn("w:fill"), hexcor)
    tcPr.append(shd)


def T(cabecalho, linhas, larguras=None, tam=9):
    """Tabela com cabeçalho escuro e linhas alternadas."""
    t = doc.add_table(rows=1, cols=len(cabecalho))
    t.style = "Table Grid"
    t.alignment = WD_TABLE_ALIGNMENT.CENTER
    for i, c in enumerate(cabecalho):
        cel = t.rows[0].cells[i]
        cel.text = ""
        r = cel.paragraphs[0].add_run(c)
        r.bold = True
        r.font.size = Pt(tam)
        r.font.color.rgb = RGBColor(0xFF, 0xF4, 0xE0)
        _sombra(cel, "3B2A1E")
    for k, linha in enumerate(linhas):
        cels = t.add_row().cells
        for i, v in enumerate(linha):
            cels[i].text = ""
            _runs(cels[i].paragraphs[0], str(v))
            for r in cels[i].paragraphs[0].runs:
                r.font.size = Pt(tam)
            if k % 2 == 1:
                _sombra(cels[i], "F3ECE2")
    if larguras:
        for linha in t.rows:
            for i, w in enumerate(larguras):
                linha.cells[i].width = Cm(w)
    doc.add_paragraph().paragraph_format.space_after = Pt(2)
    return t


def CAIXA(titulo, texto, cor="FBF1E2"):
    """Quadro de destaque (uma célula sombreada)."""
    t = doc.add_table(rows=1, cols=1)
    t.style = "Table Grid"
    cel = t.rows[0].cells[0]
    _sombra(cel, cor)
    cel.text = ""
    r = cel.paragraphs[0].add_run(titulo)
    r.bold = True
    r.font.color.rgb = FERRUGEM
    for linha in texto if isinstance(texto, list) else [texto]:
        p = cel.add_paragraph()
        _runs(p, linha)
    doc.add_paragraph().paragraph_format.space_after = Pt(2)


def IMG(rel, largura=15.5, legenda=""):
    caminho = os.path.join(ARTE, rel)
    if not os.path.exists(caminho):
        return
    doc.add_picture(caminho, width=Cm(largura))
    doc.paragraphs[-1].alignment = WD_ALIGN_PARAGRAPH.CENTER
    if legenda:
        P(legenda, italico=True, cor=CINZA, tam=8.5, alinhar="centro")


def SUMARIO():
    p = doc.add_paragraph()
    r = p.add_run()
    for tipo, txt in (("begin", None), (None, 'TOC \\o "1-2" \\h \\z \\u'), ("separate", None), (None, None), ("end", None)):
        if tipo:
            fc = OxmlElement("w:fldChar")
            fc.set(qn("w:fldCharType"), tipo)
            r._r.append(fc)
        elif txt:
            it = OxmlElement("w:instrText")
            it.set(qn("xml:space"), "preserve")
            it.text = txt
            r._r.append(it)
        else:
            t = OxmlElement("w:t")
            t.text = "Clique com o botão direito aqui e escolha \"Atualizar campo\" para montar o sumário."
            r._r.append(t)


def RODAPE(texto):
    for s in doc.sections:
        p = s.footer.paragraphs[0]
        p.text = texto
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        for r in p.runs:
            r.font.size = Pt(8)
            r.font.color.rgb = CINZA
